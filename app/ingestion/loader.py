"""
Document loaders for various file formats.

WHAT THIS DOES:
- Reads files from disk
- Extracts text content
- Returns LangChain Document objects

SUPPORTED FORMATS:
- .txt (plain text)
- .md (markdown)
- .pdf (PDF documents)
- .docx (Word documents)
"""

from pathlib import Path
from typing import List

from langchain_core.documents import Document
import structlog

logger = structlog.get_logger()


def load_text_file(file_path: str) -> List[Document]:
    """
    Load a plain text or markdown file.

    Args:
        file_path: Path to the .txt or .md file

    Returns:
        List containing one Document with the file content
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    content = path.read_text(encoding="utf-8")

    doc = Document(
        page_content=content,
        metadata={
            "source": str(path.name),
            "file_path": str(path.absolute()),
            "file_type": path.suffix.lower(),
        }
    )

    logger.info("Loaded text file", file=path.name, chars=len(content))

    return [doc]


def load_pdf_file(file_path: str) -> List[Document]:
    """
    Load a PDF file.

    Args:
        file_path: Path to the .pdf file

    Returns:
        List of Documents (one per page)
    """
    from pypdf import PdfReader

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    reader = PdfReader(str(path))
    documents = []

    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text.strip():  # Only add non-empty pages
            doc = Document(
                page_content=text,
                metadata={
                    "source": str(path.name),
                    "file_path": str(path.absolute()),
                    "file_type": ".pdf",
                    "page": i + 1,
                }
            )
            documents.append(doc)

    logger.info("Loaded PDF file", file=path.name, pages=len(documents))

    return documents


def load_docx_file(file_path: str) -> List[Document]:
    """
    Load a Word document (.docx).

    Args:
        file_path: Path to the .docx file

    Returns:
        List containing one Document with the full content
    """
    from docx import Document as DocxDocument

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    docx = DocxDocument(str(path))

    # Extract text from all paragraphs
    content = "\n".join([para.text for para in docx.paragraphs if para.text.strip()])

    doc = Document(
        page_content=content,
        metadata={
            "source": str(path.name),
            "file_path": str(path.absolute()),
            "file_type": ".docx",
        }
    )

    logger.info("Loaded DOCX file", file=path.name, chars=len(content))

    return [doc]


def load_document(file_path: str) -> List[Document]:
    """
    Load any supported document type.

    Automatically detects file type and uses appropriate loader.

    Args:
        file_path: Path to the document

    Returns:
        List of Documents

    Usage:
        docs = load_document("policy.pdf")
        docs = load_document("guide.txt")
    """
    path = Path(file_path)
    suffix = path.suffix.lower()

    loaders = {
        ".txt": load_text_file,
        ".md": load_text_file,
        ".pdf": load_pdf_file,
        ".docx": load_docx_file,
    }

    if suffix not in loaders:
        raise ValueError(f"Unsupported file type: {suffix}. Supported: {list(loaders.keys())}")

    return loaders[suffix](file_path)